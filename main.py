import docx
import openpyxl


def parse_doc(doc):
    data = []
    for table in doc.tables:
        i = 1
        stroka = []
        for row in table.rows:
            if i <= 12:
                # print('i = ', i)
                # print("Нужные данные2", row.cells[3].text)
                stroka.append(row.cells[3].text)
                i += 1
            else:
                print("stroka = ", stroka)
                i = 1
                data.append(stroka)
                stroka = []
                stroka.append(row.cells[3].text)
                i += 1
    print(data)
    return data

def make_excel(data):
    table = openpyxl.Workbook()
    sheet = table.active
    sheet.append(('Вид объекта недвижимости',
                              'Кадастровый номер',
                              'Назначение объекта недвижимости',
                              'Виды разрешенного использования объекта недвижимости',
                              'Адрес',
                              'Площадь',
                              'Вид права, доля в праве',
                              'дата государственной регистрации',
                              'номер государственной регистрации',
                              'основание государственной регистрации',
                              'дата государственной регистрации прекращения права',
                              'Ограничение прав и обременение объекта недвижимости'))
    for ctx in data:
        print("ctx", ctx)
        # Добавляем в результирующую таблицу и сохраняем ее
        sheet.append(ctx)
        table.save('D:/ProjectPython/KUMI/3.xlsx')

if __name__ == '__main__':
    path = 'D:/ProjectPython/KUMI/2.docx'
    doc = docx.Document(path)
    data = parse_doc(doc)
    make_excel(data)
